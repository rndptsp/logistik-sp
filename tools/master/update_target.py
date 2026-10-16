#!/usr/bin/env python3
"""Update master TARGET (01_TARGET.xlsx) per bulan baru.

Pemakaian (dari folder ini):
  python3 update_target.py init                       # sekali: baca 01_TARGET.xlsx -> data/target_*.csv.gz
  python3 update_target.py target --month 11 --tactical "Tactical demand SPC 2026.xlsx" \
        --snop SNOP.xlsx --monitoring "Monitoring Performance ... November 2026.xlsx" [--dry-run]
  python3 update_target.py views                      # tulis ulang 01_TARGET.xlsx dari data

Aturan bulan baru (>= Juli):
  - Volume = Tactical (Target per baris), harian = vektor harian Tactical.
  - Plant = porsi SNOP per (bulan, kode distrik, term) + asumsi plant (config/settings.json: tanpa_plant).
  - FRC: ekspeditur = proporsi sheet 'Target Ekspeditur' di file monitoring bulan itu (tanpa aturan tambahan),
    kecuali PP Belawan (Sumut): ekspeditur = distributor (sister company). CP Baturaja SMBR ikut aturan umum
    (distributor tidak mendapat target franco; konfirmasi Rendy 16-10-2026).
    Distrik tidak ada di monitoring -> proporsi bulan terdekat sebelumnya di data tersimpan, jika tidak ada TANPA_EKSPEDITUR.
  - FRC CONTAINER = SUMA 100%, FRC KAPAL = MKTC 100%; FOT/FOT CONTAINER/FOB/FAS: ekspeditur = distributor.
Bulan yang sudah ada diganti (bukan ditambah). Backup otomatis ke data/backup/.
"""
import argparse,csv,gzip,json,re,shutil,datetime as dt,collections as C
from pathlib import Path
import openpyxl

import os
BASE=Path(os.environ.get('MASTER_HOME') or Path(__file__).resolve().parent)   # data/config live outside the repo (MASTER_HOME)
CFG=BASE/'config';DATA=BASE/'data';BAK=DATA/'backup'
SET=json.load(open(CFG/'settings.json'))
MN={'Januari':1,'Februari':2,'Maret':3,'April':4,'Mei':5,'Juni':6,'Juli':7,'Agustus':8,'September':9,'Oktober':10,'November':11,'Desember':12}
INV={v:k for k,v in MN.items()}
PD=SET['target_alias_plant'];NM=SET['prov_alias'];TP=SET['tanpa_plant']
PNAME=SET['plants_snop_nama']
D31=['D%d'%i for i in range(1,32)]
TEC=['TAHUN','BULAN_NO','BULAN_NAMA','SOURCE_TYPE','INCOTERM','PROVINSI','KODE_DISTRIK','DISTRIK','KODE_PLANT','SOURCE_PLANT','PLANT_SNOP','EXPEDITUR_KODE','TARGET_TON_BULAN','PROPORSI_PCT','PROPORSI_SUMBER','CATATAN']+D31
TEC=TEC+['KEMASAN']   # kolom terakhir (setelah D31): ZAK / CURAH / JUMBO BAG
DAILY=slice(16,47)   # D1..D31 di baris TE
TDC=['SUMBER','TAHUN','BULAN_NO','BULAN','PROVINSI','KODE_DISTRIK','DISTRIK','KODE_PLANT','PLANT','TERM','TERM_ASAL','HANDLING','KETERANGAN','BRAND','DISTRIBUTOR','SOLD_TO','TARGET_TON','PLANT_SOURCE','CATATAN']+D31
NUMC={'TAHUN','BULAN_NO','KODE_DISTRIK','TARGET_TON_BULAN','PROPORSI_PCT','TARGET_TON'}|set(D31)
nz=lambda x:str(x).replace(' ','').upper()
def log(*a): print(*a,flush=True)

def sv(v):
    if v is None: return ''
    if isinstance(v,float): return repr(round(v,6))
    return str(v)
def save(path,cols,rows):
    with gzip.open(path,'wt',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(cols)
        for r in rows: w.writerow([sv(x) for x in r])
def load(path,cols):
    out=[]
    with gzip.open(path,'rt',newline='',encoding='utf-8') as f:
        rd=csv.reader(f);h=next(rd);assert h==cols or h==[c for c in cols if c!='KEMASAN'],'kolom tidak cocok'
        for r in rd:
            o=[]
            if len(r)<len(cols) and 'KEMASAN' in cols: r=r+['ZAK']   # berkas lama tanpa KEMASAN = zak
            for c,x in zip(cols,r):
                if x=='': o.append(None)
                elif c in NUMC:
                    v=float(x);o.append(int(v) if c in ('TAHUN','BULAN_NO','KODE_DISTRIK') and v==int(v) else v)
                else: o.append(x)
            out.append(o)
    return out
def backup(p):
    BAK.mkdir(exist_ok=True)
    if p.exists(): shutil.copy2(p,BAK/(p.name+'.'+dt.datetime.now().strftime('%Y%m%d_%H%M%S')))

# ---------------------------------------------------------------- init
def cmd_init(a):
    p=BASE/'01_TARGET.xlsx'
    if (DATA/'target_ekspeditur.csv.gz').exists() and not a.force: raise SystemExit('Sudah ada; pakai --force untuk menimpa.')
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
    for sh,cols,fn in (('Target Ekspeditur',TEC,'target_ekspeditur'),('Target Distributor',TDC,'target_distributor')):
        it=wb[sh].iter_rows(values_only=True);h=list(next(it));assert h==cols or h==[c for c in cols if c!='KEMASAN'],(sh,h[:5])
        rows=[list(r) for r in it if r[0] is not None]
        if 'KEMASAN' in cols and 'KEMASAN' not in h: rows=[r+['ZAK'] for r in rows]
        for r in rows:
            for c in ('KODE_PLANT',):
                i=cols.index(c)
                if r[i] is not None: r[i]=str(r[i])
        save(DATA/(fn+'.csv.gz'),cols,rows);log(sh,len(rows),'baris ->',fn+'.csv.gz')
    pan=[[c for c in r] for r in wb['PANDUAN'].iter_rows(values_only=True)]
    json.dump(pan,open(CFG/'panduan_target.json','w'),ensure_ascii=False)

# ---------------------------------------------------------------- reading inputs
def read_tactical(path,mname):
    ws=openpyxl.load_workbook(path,read_only=True,data_only=True)['Sheet1']
    it=ws.iter_rows(values_only=True);next(it);h=list(next(it));ix={k:i for i,k in enumerate(h)}
    need=['PROP','BRAND','Segmen','Sdis','Daerah','Sold to','DISTRIBUTOR','SYARAT','Target','Periode']
    for n in need: assert n in ix,'kolom Tactical hilang: '+n
    rows=[]
    for r in it:
        if r[ix['Periode']]!=mname or r[ix['SYARAT']] is None: continue
        t=r[ix['Target']]
        if not isinstance(t,(int,float)): continue
        d={n:r[ix[n]] for n in need};d['D']=[(r[ix[i]] if isinstance(r[ix[i]],(int,float)) else 0.0) for i in range(1,32)] if 1 in ix else [0.0]*31
        rows.append(d)
    return rows
def read_snop(path):
    ws=openpyxl.load_workbook(path,read_only=True,data_only=True)['Rincian']
    it=ws.iter_rows(values_only=True);h=list(next(it));ix={k:i for i,k in enumerate(h)}
    out=[]
    for r in it:
        if r[ix['KEMASAN']] is None: continue
        if str(r[ix['KEMASAN']]).upper()!='ZAK' or str(r[ix['KETERANGAN']]).lower()=='corsales': continue
        out.append(dict(per=r[ix['Periode']],kd=int(r[ix['KODE DISTRIK']]),term=r[ix['TERM']],kp=str(r[ix['KODE DELIV. PLANT']]),pn=r[ix['DELIV. PLANT']],h=r[ix['HANDLING']],ket=str(r[ix['KETERANGAN']]).lower(),vol=r[ix['Volume']] or 0,prov=r[ix['PROVINSI']]))
    return out
def read_monitoring(path,mn):
    ws=openpyxl.load_workbook(path,read_only=True,data_only=True)['Target Ekspeditur']
    mon={}
    for r in list(ws.iter_rows(values_only=True))[2:]:
        if not r[1] or not r[3] or not r[4]: continue
        v=sum((x or 0) for x in r[6:37] if isinstance(x,(int,float)))
        e='SILOG' if r[4]=='SILC' else r[4]
        d=mon.setdefault(nz(r[3]),{});d[e]=d.get(e,0)+v
    return mon

# ---------------------------------------------------------------- build month
def build(mn,tact,snop,mon,te_old,td_old):
    mname=INV[mn];code={}
    sh_n=C.defaultdict(lambda:C.defaultdict(float));sh_nr=C.defaultdict(lambda:C.defaultdict(float));sh_r=C.defaultdict(lambda:C.defaultdict(float));sh_a=C.defaultdict(lambda:C.defaultdict(float));sh_h=C.defaultdict(lambda:C.defaultdict(float));psh=C.defaultdict(lambda:C.defaultdict(float))
    for r in snop:
        code[r['pn']]=r['kp']
        if r['term']=='FOT': psh[(r['kd'],'FOT CONTAINER' if r['h']=='CONTAINER' else 'FOT')][r['pn']]+=r['vol']
        if r['per']!=mname: continue
        k=(r['kd'],r['term']);pk=(r['kp'],r['pn'])
        sh_a[k][pk]+=r['vol']
        if r['ket']=='retail': sh_r[k][pk]+=r['vol']
        if r['h']=='NONE':
            sh_n[k][pk]+=r['vol']
            if r['ket']=='retail': sh_nr[k][pk]+=r['vol']
        sh_h[k+(r['h'],)][pk]+=r['vol']
    for n,k in PNAME.items(): code.setdefault(k,n)
    ok=lambda s:s is not None and sum(s.values())>0
    td=[]
    for r in tact:
        tb=r['SYARAT'].split()[0];hd='CONTAINER' if 'CONTAINER' in r['SYARAT'] else ('POLYSLING' if 'KAPAL' in r['SYARAT'] else None)
        k=(int(r['Sdis']),tb)
        if hd and ok(sh_h.get(k+(hd,))): sh=sh_h[k+(hd,)]
        elif not hd and ok(sh_nr.get(k)): sh=sh_nr[k]
        elif not hd and ok(sh_n.get(k)): sh=sh_n[k]
        elif ok(sh_r.get(k)): sh=sh_r[k]
        elif ok(sh_a.get(k)): sh=sh_a[k]
        else: sh=None
        p=NM.get(r['PROP'],r['PROP'])
        base=['TACTICAL',2026,mn,mname,p,int(r['Sdis']),r['Daerah']]
        mid=[tb,r['SYARAT'],hd or 'NONE',r['Segmen'],r['BRAND'],r['DISTRIBUTOR'],r['Sold to']]
        if sh is None:
            td.append(base+[None,None]+mid+[r['Target'],'TANPA_PLANT','distrik-term tidak ada di SNOP bulan itu']+list(r['D']));continue
        tot=sum(sh.values())
        for (kp,pn),v in sh.items():
            f=v/tot
            if f==0: continue
            td.append(base+[kp,pn]+mid+[r['Target']*f,'SNOP_TUNGGAL' if len(sh)==1 else 'SNOP_PORSI_%.4f'%f,None]+[x*f for x in r['D']])
    # daily weights
    W={}
    for sy in ('FRC','FRC CONTAINER','FRC KAPAL'):
        rs=[r for r in tact if r['SYARAT']==sy]
        if rs:
            tot=[sum(r['D'][i] for r in rs) for i in range(31)];T=sum(tot)
            if T>0: W[sy]=[x/T for x in tot]
    assert 'FRC' in W,'tidak ada vektor harian FRC di Tactical'
    # distributor shares (bulan ini) per (kd,inc)
    dist=C.defaultdict(lambda:C.defaultdict(float))
    for r in tact:
        if r['Target']>0: dist[(int(r['Sdis']),r['SYARAT'])][r['DISTRIBUTOR']]+=r['Target']
    # stored FRC proportions for fallback
    old=C.defaultdict(lambda:C.defaultdict(lambda:C.defaultdict(float)))
    ti={c:i for i,c in enumerate(TEC)}
    for r in te_old:
        if r[ti['INCOTERM']]=='FRC' and r[ti['BULAN_NO']]<mn and r[ti['SOURCE_TYPE']]=='TACTICAL' and r[ti['EXPEDITUR_KODE']] not in ('TANPA_EKSPEDITUR',) and r[ti['KEMASAN']] in (None,'ZAK'):
            old[(r[ti['PROVINSI']],nz(r[ti['DISTRIK']]))][r[ti['BULAN_NO']]][r[ti['EXPEDITUR_KODE']]]+=r[ti['TARGET_TON_BULAN']] or 0
    # aggregate TD -> TE keys
    x={c:i for i,c in enumerate(TDC)}
    frc=C.defaultdict(float);oth={}
    for r in td:
        inc=r[x['TERM_ASAL']];kp=r[x['KODE_PLANT']];pn=r[x['PLANT']] or 'TANPA_PLANT'
        if inc in ('FRC','FRC CONTAINER','FRC KAPAL'): frc[(r[x['PROVINSI']],r[x['KODE_DISTRIK']],r[x['DISTRIK']],kp,pn,inc)]+=r[x['TARGET_TON']]
        else:
            k=(r[x['PROVINSI']],r[x['KODE_DISTRIK']],r[x['DISTRIK']],kp,pn,inc,r[x['DISTRIBUTOR']],r[x['SOLD_TO']])
            a=oth.setdefault(k,[0.0]+[0.0]*31);a[0]+=r[x['TARGET_TON']]
            for i in range(31): a[i+1]+=r[x['D1']+i] or 0
    out=[];stat=C.Counter();tanpa=C.defaultdict(float)
    def add(p,kd,d,kp,pn,e,v,src,cat,inc,w):
        dl=[v*y for y in w] if w else [0.0]*31
        out.append([2026,mn,mname,'TACTICAL',inc,p,kd,d,kp,PD.get(pn,pn),pn,e,v,1.0 if v else 0.0,src,cat]+dl)
    for (p,kd,d,kp,pn,inc),v in sorted(frc.items(),key=str):
        if inc!='FRC':
            add(p,kd,d,kp,pn,{'FRC CONTAINER':'SUMA','FRC KAPAL':'MKTC'}[inc],v,'TUNGGAL_100','ekspeditur tunggal 100% (konfirmasi Rendy)',inc,W.get(inc) or W['FRC']);continue
        w=W['FRC']
        def emit(sh_,src,cat):
            tot_=sum(sh_.values())
            for e_,x_ in sh_.items():
                if x_<=0: continue
                f_=x_/tot_;vv=v*f_
                out.append([2026,mn,mname,'TACTICAL','FRC',p,kd,d,kp,PD.get(pn,pn),pn,e_,vv,f_,src,cat]+[vv*y for y in w])
        if p=='SUMATERA UTARA' and 'Belawan' in str(pn):
            z=dist.get((kd,'FRC'))
            if z: emit(dict(z),'DISTRIBUTOR_TACTICAL_SUMUT','Belawan: ekspeditur = distributor (sister company, Rendy)');continue
        m_=mon.get(nz(d))
        if m_ and sum(m_.values())>0: emit(m_,'MONITORING_%d'%mn,None);continue
        o_=old.get((p,nz(d)))
        if o_:
            bm=max(o_);emit(dict(o_[bm]),'FALLBACK_BULAN_%d'%bm,'distrik tidak ada di monitoring; proporsi bulan %d'%bm);stat['fallback']+=1;continue
        stat['tanpa']+=1;tanpa[(p,d)]+=v;emit({'TANPA_EKSPEDITUR':1.0},'TANPA','tidak ada proporsi (monitoring/bulan lalu)')
    for (p,kd,d,kp,pn,inc,dn,so),a in sorted(oth.items(),key=str):
        out.append([2026,mn,mname,'TACTICAL',inc,p,kd,d,kp,PD.get(pn,pn),pn,dn,a[0],None,'DISTRIBUTOR_TACTICAL','SOLD_TO %s'%so]+a[1:])
    # plant assumptions
    ti={c:i for i,c in enumerate(TEC)}
    def flag(r,txt): r[15]=((r[15]+'; ') if r[15] else '')+txt
    def setp(r,pn): r[8]=code.get(pn,code.get(pn));r[9]=PD.get(pn,pn);r[10]=pn
    for r in out:
        if r[10]=='TANPA_PLANT':
            inc=r[4]
            if inc=='FOB': setp(r,PNAME[TP['FOB']]);flag(r,'PLANT ASUMSI (Teluk Bayur untuk FOB)')
            elif inc=='FRC': setp(r,PNAME[TP['FRC']]);flag(r,'PLANT ASUMSI (Indarung untuk FRC)')
        if r[6] in (101004,) and r[4]=='FOT' and r[10]=='CP Indarung SP':
            r[8]=None;r[9]='TANPA_PLANT';r[10]='TANPA_PLANT';flag(r,'plant Indarung dibuang (FOT Singkil bukan dari Indarung)')
    new=[];left=C.Counter();pstat=C.Counter()
    for r in out:
        if r[10]=='TANPA_PLANT' and r[4] in ('FOT','FOT CONTAINER'):
            sh_=None if r[6]==101004 else psh.get((r[6],r[4]))
            if r[5]=='SUMATERA UTARA': sh_={PNAME[TP['FOT_sumut']]:1.0}
            if sh_ and sum(sh_.values())>0:
                tt=sum(sh_.values())
                for pn,v in sh_.items():
                    f=v/tt;n=list(r);setp(n,pn);n[12]=r[12]*f;n[16:]=[y*f for y in r[16:]]
                    flag(n,'PLANT ASUMSI porsi SNOP %.0f%%'%(f*100));new.append(n);pstat[pn]+=n[12]
                continue
            if r[4]=='FOT CONTAINER' or r[6]==101004:
                setp(r,PNAME[TP['FOT CONTAINER']]);flag(r,'PLANT ASUMSI (Teluk Bayur)');pstat['Teluk Bayur']+=r[12]
            else: left[r[5]]+=r[12]
        new.append(r)
    return td,new,dict(stat),{k:v for k,v in tanpa.items()},dict(left),dict(pstat),W

def checks(mn,tact,td,te):
    tt=C.defaultdict(float)
    for r in tact: tt[r['SYARAT']]+=r['Target']
    x={c:i for i,c in enumerate(TDC)};ti={c:i for i,c in enumerate(TEC)}
    msgs=[];bad=0
    td_t=C.defaultdict(float);te_t=C.defaultdict(float);te_d=C.defaultdict(float)
    for r in td: td_t[r[x['TERM_ASAL']]]+=r[x['TARGET_TON']]
    for r in te:
        te_t[r[ti['INCOTERM']]]+=r[ti['TARGET_TON_BULAN']];te_d[r[ti['INCOTERM']]]+=sum(r[DAILY])
    for inc in sorted(tt):
        a,b,c,d=tt[inc],td_t[inc],te_t[inc],te_d[inc]
        flag='OK' if abs(a-b)<0.01 and abs(a-c)<0.01 and abs(c-d)<0.5 else 'SELISIH'
        if flag!='OK': bad+=1
        msgs.append('  %-14s Tactical %12.2f | TD %12.2f | TE %12.2f | Σharian %12.2f  %s'%(inc,a,b,c,d,flag))
    return msgs,bad

# ---------------------------------------------------------------- commands

# ---------------------------------------------------------------- curah & jumbo (SNOP korsales)
KEM_SNOP={'ZAK':'ZAK','CURAH':'CURAH','JUMBO':'JUMBO BAG'}
def read_snop_nonzak(path):
    """Semua baris SNOP selain zak (CURAH, JUMBO; di SNOP semuanya korsales). Tidak ada distributor/ekspeditur/harian di SNOP."""
    ws=openpyxl.load_workbook(path,read_only=True,data_only=True)['Rincian']
    it=ws.iter_rows(values_only=True);h=list(next(it));ix={k:i for i,k in enumerate(h)};out=[]
    for r in it:
        if r[ix['KEMASAN']] is None: continue
        k=KEM_SNOP.get(str(r[ix['KEMASAN']]).upper(),str(r[ix['KEMASAN']]).upper())
        if k=='ZAK' or not (r[ix['Volume']] or 0): continue
        out.append(dict(per=r[ix['Periode']],kd=int(r[ix['KODE DISTRIK']]),distrik=r[ix['SALES DISTRIK']],term=r[ix['TERM']],kp=str(r[ix['KODE DELIV. PLANT']]),pn=r[ix['DELIV. PLANT']],h=r[ix['HANDLING']],vol=r[ix['Volume']],prov=NM.get(r[ix['PROVINSI']],r[ix['PROVINSI']]),kem=k))
    return out
def snop_inc(term,h):
    if term=='FRC': return {'NONE':'FRC','CONTAINER':'FRC CONTAINER','POLYSLING':'FRC KAPAL'}.get(h,'FRC')
    if term=='FOT': return 'FOT CONTAINER' if h=='CONTAINER' else 'FOT'
    return term
def build_nonzak(mn,rows,te_ref):
    """Baris TE untuk curah/jumbo bulan mn: volume SNOP apa adanya, plant dari SNOP, harian = rata per hari,
    ekspeditur/distributor tidak diketahui (TANPA_EKSPEDITUR / TANPA_DISTRIBUTOR). Provinsi non-kinerja dilewati."""
    import calendar
    nk=set(SET.get('provinsi_non_kinerja',[]));ti={c:i for i,c in enumerate(TEC)}
    spell={}
    for r in te_ref: spell.setdefault(nz(r[ti['DISTRIK']]),r[ti['DISTRIK']])
    nd=calendar.monthrange(2026,mn)[1];agg=C.defaultdict(float)
    for r in rows:
        if r['per']!=INV[mn] or r['prov'] in nk: continue
        d=spell.get(nz(r['distrik']),re.sub(r'^(KAB|KOTA)\.(?=\S)',r'\1. ',r['distrik'].strip()))
        agg[(r['prov'],r['kd'],d,r['kp'],r['pn'],snop_inc(r['term'],r['h']),r['kem'])]+=r['vol']
    out=[]
    for (p,kd,d,kp,pn,inc,kem),v in sorted(agg.items(),key=str):
        e='TANPA_EKSPEDITUR' if inc.startswith('FRC') else 'TANPA_DISTRIBUTOR'
        dl=[v/nd if i<nd else 0.0 for i in range(31)]
        out.append([2026,mn,INV[mn],'SNOP',inc,p,kd,d,kp,PD.get(pn,pn),pn,e,v,1.0,'SNOP_KORSALES','%s: volume SNOP, tanpa %s, harian rata'%(kem,'ekspeditur' if inc.startswith('FRC') else 'distributor')]+dl+[kem])
    return out
def check_nonzak(mn,rows,out):
    ti={c:i for i,c in enumerate(TEC)};nk=set(SET.get('provinsi_non_kinerja',[]))
    A=C.defaultdict(float);B=C.defaultdict(float)
    for r in rows:
        if r['per']==INV[mn] and r['prov'] not in nk: A[(snop_inc(r['term'],r['h']),r['pn'],r['kem'])]+=r['vol']
    for r in out: B[(r[ti['INCOTERM']],r[ti['PLANT_SNOP']],r[ti['KEMASAN']])]+=r[ti['TARGET_TON_BULAN']]
    bad=[(k,A.get(k,0),B.get(k,0)) for k in set(A)|set(B) if abs(A.get(k,0)-B.get(k,0))>0.01]
    return A,bad
def cmd_nonzak(a):
    """Isi/isi ulang target CURAH & JUMBO dari SNOP untuk bulan --months (mis. 1-10); baris ZAK tidak disentuh."""
    lo,hi=(int(x) for x in a.months.split('-')) if '-' in a.months else (int(a.months),int(a.months))
    te=load(DATA/'target_ekspeditur.csv.gz',TEC);ti={c:i for i,c in enumerate(TEC)}
    rows=read_snop_nonzak(a.snop);log('SNOP non-zak baris:',len(rows))
    keep=[r for r in te if not (lo<=r[ti['BULAN_NO']]<=hi and r[ti['KEMASAN']] not in (None,'ZAK'))]
    new=[];tot=C.defaultdict(float);bad_all=[]
    for mn in range(lo,hi+1):
        o=build_nonzak(mn,rows,[r for r in keep if r[ti['KEMASAN']] in (None,'ZAK')]);A,bad=check_nonzak(mn,rows,o);new+=o;bad_all+=bad
        for r in o: tot[(mn,r[ti['KEMASAN']])]+=r[ti['TARGET_TON_BULAN']]
    log('Target non-zak per bulan×kemasan (ton):',{'%d %s'%k:round(v) for k,v in sorted(tot.items())})
    log('Selisih vs SNOP per bulan×incoterm×plant×kemasan:',len(bad_all),'kombinasi tidak cocok',bad_all[:5])
    if bad_all and not a.force: raise SystemExit('REKONSILIASI GAGAL.')
    if a.dry_run: log('DRY-RUN.');return
    p=DATA/'target_ekspeditur.csv.gz';backup(p);allr=keep+new;allr.sort(key=lambda r:(r[ti['BULAN_NO']],))
    save(p,TEC,allr);log('target_ekspeditur ditulis:',len(allr),'baris (non-zak baru: %d)'%len(new))

def cmd_target(a):
    mn=a.month
    if mn<7: raise SystemExit('Skrip ini untuk bulan >= 7 (Juli ke atas).')
    for f in (a.tactical,a.snop,a.monitoring):
        if not Path(f).exists(): raise SystemExit('File tidak ada: '+f)
    te_old=load(DATA/'target_ekspeditur.csv.gz',TEC);td_old=load(DATA/'target_distributor.csv.gz',TDC)
    tact=read_tactical(a.tactical,INV[mn]);log('Tactical baris bulan %s: %d'%(INV[mn],len(tact)))
    if not tact: raise SystemExit('Tidak ada baris Tactical untuk bulan itu (kolom Periode).')
    snop=read_snop(a.snop);log('SNOP baris (ZAK non-corsales):',len(snop))
    mon=read_monitoring(a.monitoring,mn);log('Monitoring: distrik dengan proporsi =',len(mon))
    td,te,stat,tanpa,left,pstat,W=build(mn,tact,snop,mon,te_old,td_old)
    for r in te: r.append('ZAK')
    msgs,bad=checks(mn,tact,td,te)
    log('Rekonsiliasi ZAK (ton):');[log(m) for m in msgs]
    nzr=read_snop_nonzak(a.snop);nzo=build_nonzak(mn,nzr,[r for r in te_old if r[len(TEC)-1] in (None,'ZAK')]+te);_,nzbad=check_nonzak(mn,nzr,nzo)
    log('Curah/jumbo (SNOP): %d baris, %.1f t | selisih vs SNOP: %d kombinasi'%(len(nzo),sum(r[12] for r in nzo),len(nzbad)))
    if nzbad: bad+=1
    te=te+nzo
    log('Fallback FRC:',stat.get('fallback',0),'| TANPA_EKSPEDITUR:',stat.get('tanpa',0),{str(k):round(v,1) for k,v in tanpa.items()})
    log('Plant asumsi (ton):',{k:round(v,1) for k,v in pstat.items()})
    if left: log('PERINGATAN FOT tanpa plant tersisa:',{k:round(v,1) for k,v in left.items()})
    ti={c:i for i,c in enumerate(TEC)}
    # perbandingan dengan data tersimpan (bulan sama) bila ada
    ex=[r for r in te_old if r[ti['BULAN_NO']]==mn and r[ti['KEMASAN']] in (None,'ZAK')]
    if ex:
        A=C.defaultdict(float);B=C.defaultdict(float)
        for r in ex: A[(r[ti['INCOTERM']],r[ti['EXPEDITUR_KODE']])]+=r[ti['TARGET_TON_BULAN']] or 0
        for r in te: B[(r[ti['INCOTERM']],r[ti['EXPEDITUR_KODE']])]+=r[ti['TARGET_TON_BULAN']]
        diff=[(k,round(B.get(k,0)-A.get(k,0),2)) for k in sorted(set(A)|set(B),key=str) if abs(B.get(k,0)-A.get(k,0))>0.5]
        log('Beda vs data tersimpan bulan ini (ton, baru-lama) pada incoterm×ekspeditur:',len(diff),'kombinasi',diff[:12])
    if bad and not a.force: raise SystemExit('REKONSILIASI GAGAL - tidak ditulis (gunakan --force bila sengaja).')
    if a.dry_run: log('DRY-RUN: tidak ada yang ditulis.');return
    for fn,cols,old,new,idx in (('target_ekspeditur',TEC,te_old,te,ti['BULAN_NO']),('target_distributor',TDC,td_old,td,TDC.index('BULAN_NO'))):
        p=DATA/(fn+'.csv.gz');backup(p)
        keep=[r for r in old if r[idx]!=mn];allr=keep+new
        allr.sort(key=lambda r:(r[idx],))
        save(p,cols,allr);log(fn,'ditulis:',len(allr),'baris (bulan %d: %d baris)'%(mn,len(new)))
    log('Selanjutnya: python3 update_target.py views')

def cmd_views(a):
    import xlsxwriter
    te=load(DATA/'target_ekspeditur.csv.gz',TEC);td=load(DATA/'target_distributor.csv.gz',TDC)
    p=BASE/'01_TARGET.xlsx'
    if p.exists(): backup(p)
    wb=xlsxwriter.Workbook(str(p),{'nan_inf_to_errors':True});hf=wb.add_format({'bold':True,'font_color':'white','bg_color':'#3B4A63'})
    for name,cols,rows in (('Target Ekspeditur',TEC,te),('Target Distributor',TDC,td)):
        ws=wb.add_worksheet(name);ws.write_row(0,0,cols,hf)
        for i,r in enumerate(rows,1):
            for j,v in enumerate(r):
                if v is not None: ws.write(i,j,v)
        ws.freeze_panes(1,0);ws.autofilter(0,0,len(rows),len(cols)-1)
    ws=wb.add_worksheet('CEK');ws.write_row(0,0,['Bulan','Incoterm','Target Ekspeditur (ton)','Target Distributor (ton)','Selisih','Σ D1..D31 - total'],hf)
    ti={c:i for i,c in enumerate(TEC)};x={c:i for i,c in enumerate(TDC)}
    A=C.defaultdict(float);B=C.defaultdict(float);Dd=C.defaultdict(float)
    for r in te:
        k=(r[1],r[4]);A[k]+=r[12] or 0;Dd[k]+=sum((v or 0) for v in r[DAILY])
    for r in td: B[(r[2],r[10])]+=r[16] or 0
    i=1
    for k in sorted(A):
        ws.write_row(i,0,[k[0],k[1],round(A[k],2),round(B.get(k,0),2),round(A[k]-B.get(k,0),2),round(Dd[k]-A[k],4)]);i+=1
    ws.write(i+1,0,'Catatan: selisih Ekspeditur-Distributor ≠ 0 pada Jan–Mar (volume SNOP) dan baris MANUAL (PSK) wajar.')
    ws=wb.add_worksheet('PANDUAN')
    try:
        for i,r in enumerate(json.load(open(CFG/'panduan_target.json'))): ws.write_row(i,0,[v for v in r if v is not None])
    except FileNotFoundError: ws.write(0,0,'Lihat docstring update_target.py')
    wb.close();log('01_TARGET.xlsx ditulis (TE %d, TD %d baris). Jangan buka file saat proses.'%(len(te),len(td)))

def main():
    ap=argparse.ArgumentParser();sp=ap.add_subparsers(dest='cmd',required=True)
    s=sp.add_parser('init');s.add_argument('--force',action='store_true')
    s=sp.add_parser('target');s.add_argument('--month',type=int,required=True);s.add_argument('--tactical',required=True);s.add_argument('--snop',required=True);s.add_argument('--monitoring',required=True);s.add_argument('--dry-run',action='store_true');s.add_argument('--force',action='store_true')
    s=sp.add_parser('nonzak');s.add_argument('--months',required=True);s.add_argument('--snop',required=True);s.add_argument('--dry-run',action='store_true');s.add_argument('--force',action='store_true')
    sp.add_parser('views')
    a=ap.parse_args()
    {'init':cmd_init,'target':cmd_target,'views':cmd_views,'nonzak':cmd_nonzak}[a.cmd](a)
if __name__=='__main__': main()
