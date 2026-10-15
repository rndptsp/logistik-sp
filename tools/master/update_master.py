#!/usr/bin/env python3
"""Update master baru logistik (Rilis + SO) dengan upsert dari file kumulatif.

Pemakaian (dari folder ini):
  python3 update_master.py rilis --input "Realisasi Shipment MTD 15 Oktober 2026.xlsx"
  python3 update_master.py so    --input "Sales Order (Source CSMS All Propinsi).xlsx"
  python3 update_master.py views            # tulis ulang 02_RILIS.xlsx dan 03_SO.xlsx dari data
Opsi: --dry-run (hanya laporan, tidak menulis), --force (abaikan guardrail), --sheets 7912,79B4 (uji sebagian).
Data: data/rilis.csv.gz, data/so.csv.gz.  Konfigurasi: config/*.json.  Log: logs/.
"""
import argparse,csv,gzip,json,re,shutil,sys,statistics,datetime as dt,collections as C
from pathlib import Path
import openpyxl

import os
BASE=Path(os.environ.get('MASTER_HOME') or Path(__file__).resolve().parent)   # data/config live outside the repo (MASTER_HOME)
CFG=BASE/'config';DATA=BASE/'data';LOGS=BASE/'logs';BAK=DATA/'backup'
SET=json.load(open(CFG/'settings.json'))
EKS=json.load(open(CFG/'ekspeditur_map.json'))
SCOPE=set(SET['provinsi_scope']);PLANTS=SET['plants'];PLANT_BATAS=SET.get('provinsi_plant_batas',{});SO_DISTRIK=SET.get('so_distrik_provinsi',{});SO_KET=str(SET.get('so_keterangan','AP1')).upper();SO_BUANG=set(SET.get('so_buang_tersimpan_provinsi',[]))
TMIN=dt.datetime.strptime(SET['tanggal_min'],'%Y-%m-%d');SILOG_END=dt.datetime.strptime(SET['silog_sampai'],'%Y-%m-%d')
LOG=[]
def log(*a):
    s=' '.join(str(x) for x in a);print(s,flush=True);LOG.append(s)
def pdate(x):
    if x is None or x=='' : return None
    if isinstance(x,dt.datetime): return x
    if isinstance(x,dt.date): return dt.datetime(x.year,x.month,x.day)
    for f in ('%d-%m-%Y','%Y-%m-%d %H:%M:%S','%Y-%m-%d'):
        try: return dt.datetime.strptime(str(x).strip(),f)
        except ValueError: pass
    return None
def num(x):
    if x is None or x=='': return None
    try: return float(x)
    except (TypeError,ValueError): return None
def sv(v):  # ke string untuk CSV / perbandingan
    if v is None: return ''
    if isinstance(v,dt.datetime): return v.strftime('%Y-%m-%d %H:%M:%S') if (v.hour or v.minute or v.second) else v.strftime('%Y-%m-%d')
    if isinstance(v,float): return repr(round(v,6))
    return str(v)
def kemasan_prod(prod):
    p=str(prod or '').upper()
    if 'CURAH' in p or 'BULK' in p: return 'CURAH'
    if re.search(r'\bJB\b|JUMBO|BAG',p) and not re.search(r'\d+\s*KG',p): return 'JUMBO BAG'
    return 'ZAK'
def kemasan_uom(uom):
    return {'TO':'CURAH','BAG':'JUMBO BAG'}.get(str(uom or '').strip().upper(),'ZAK')
def norm_distrik(s):
    s=re.sub(r'\s+',' ',str(s or '').strip())
    return re.sub(r'^(KAB|KOTA)\.(?=\S)',r'\1. ',s)
def load_store(path,cols):
    st={}
    if not path.exists(): return st
    with gzip.open(path,'rt',encoding='utf-8',newline='') as f:
        for r in csv.DictReader(f):
            for c in cols: r.setdefault(c,'')   # kolom baru ditambahkan -> kosong dulu
            if 'KEMASAN' in cols and not r['KEMASAN']: r['KEMASAN']=kemasan_prod(r.get('PRODUK'))
            st[r[cols[0]]]=r
    return st
def save_store(path,cols,rows):
    BAK.mkdir(parents=True,exist_ok=True)
    if path.exists(): shutil.copy2(path,BAK/(path.stem.replace('.csv','')+dt.datetime.now().strftime('_%Y%m%d_%H%M%S')+'.csv.gz'))
    tmp=path.with_suffix('.tmp')
    with gzip.open(tmp,'wt',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=cols);w.writeheader()
        for r in rows: w.writerow(r)
    tmp.replace(path)
def header_map(h):
    return {str(x).strip():i for i,x in enumerate(h) if x is not None}

# =============== RILIS ===============
RCOLS=['NO_SPJ','TGL_SPJ','BULAN_NO','KD_PLANT','SOURCE_PLANT','INCOTERM','INCOTERM_ASAL','ROUTE','NAMA_PROPINSI','NAMA_AREA_DISTRIK','NAMA_KEC_SUBDISTRIK','KD_DISTRIBUTOR','DISTRIBUTOR','KD_EKSPEDITUR_ASAL','EKSPEDITUR_ASAL','EXPEDITUR_KODE','EXPEDITUR_NAMA','NOPOL','NO_SO','PRODUK','KEMASAN','KEMASAN_SUMBER','TONASE','KODE_TOKO','NAMA_TOKO','ALAMAT_TOKO','NAMA_KAPAL','STATUS','WAKTU_ANTRI','WAKTU_MASUK','WAKTU_KELUAR','WAITING_TIME_JAM','DWELL_TIME_JAM','DWELL_WAITING_FLAG','DWELL_WAITING_NOTE','BERAT_KOSONG','BERAT_ISI','BERAT_REAL','KODE_KANTONG','ADA_DI_MASTER_LAMA','FLAG']
R_DT=['TGL_SPJ','WAKTU_ANTRI','WAKTU_MASUK','WAKTU_KELUAR'];R_NUM=['NO_SPJ','BULAN_NO','TONASE','WAITING_TIME_JAM','DWELL_TIME_JAM','BERAT_KOSONG','BERAT_ISI','BERAT_REAL']
NEW_CODES={}
def eks_init(name):
    if name in EKS['initials']: return EKS['initials'][name]
    ws=[w for w in re.sub(r'[^A-Za-z ]',' ',name.upper()).split() if w not in ('PT','TBK','CV','UD','KOPERASI','PERSERO','THE')]
    base=''.join(w[0] for w in ws)[:5] or 'EKS';c=base;i=2;used=set(EKS['used_codes'])
    while c in used: c=base+str(i);i+=1
    EKS['used_codes'].append(c);EKS['initials'][name]=c;EKS['kamus'][c]=name;NEW_CODES[name]=c
    return c
def rilis_row(rec,code):
    g=lambda k:rec.get(k)
    tgl=pdate(g('TGL SPJ'));prod=str(g('PRODUK') or '').strip();inc0=g('INCOTERM');inc=inc0;flags=[]
    if prod.upper().endswith('CON'):
        if inc0=='FOT': inc='FOT CONTAINER'
        else: flags.append('PRODUK_CON_TETAP_%s'%inc0)
    kde=str(g('KD EXPEDITUR')) if g('KD EXPEDITUR') not in (None,'',0,'0') else '';en=g('EXPEDITUR') or '';dist=g('DISTRIBUTOR') or ''
    alias=next((v for k,v in EKS.get('alias_prefix',{}).items() if str(en).upper().startswith(k)),None)
    if str(en).upper().startswith('ANGKUT SENDIRI') or not kde: ek,eknm=dist,dist;flags.append('EKSPEDITUR=DISTRIBUTOR')
    elif kde in EKS['by_kd']: ek,eknm=EKS['by_kd'][kde]
    elif alias: ek,eknm=alias,en;flags.append('ALIAS=%s'%alias)
    elif en in EKS['by_name']: ek,eknm=EKS['by_name'][en],en
    else:
        ek=eks_init(en);eknm=en;flags.append('KODE_EKSPEDITUR_BARU')
    ta=None
    def tm(t,j):
        d=pdate(t)
        if d is None or j in (None,''): return None
        try: hh=dt.datetime.strptime(str(j),'%H:%M:%S')
        except ValueError: return None
        return d.replace(hour=hh.hour,minute=hh.minute,second=hh.second)
    ta=tm(g('TGL ANTRI'),g('JAM ANTRI'));tmasuk=tm(g('TGL MASUK'),g('JAM MASUK'));tk=tm(g('TGL KELUAR'),g('JAM KELUAR'));note=None
    if ta is None and tmasuk is not None: ta=tmasuk;note='ANTRI disamakan dengan MASUK (self-haul/FOT Container)'
    wt=round((tmasuk-ta).total_seconds()/3600,3) if ta and tmasuk else None
    dw=round((tk-tmasuk).total_seconds()/3600,3) if tk and tmasuk else None
    fl=None
    if wt is not None and wt<0: fl='CHECK_NEGATIVE_WAITING'
    elif (wt is not None and wt>48) or (dw is not None and dw>48): fl='CHECK_OVER_48H'
    spj=str(g('NO SPJ')).strip()
    row={'NO_SPJ':spj,'TGL_SPJ':tgl,'BULAN_NO':tgl.month if tgl else None,'KD_PLANT':code,'SOURCE_PLANT':PLANTS.get(code,code),'INCOTERM':inc,'INCOTERM_ASAL':inc0,'ROUTE':g('ROUTE'),'NAMA_PROPINSI':g('NAMA PROPINSI'),'NAMA_AREA_DISTRIK':norm_distrik(g('NAMA AREA')),'NAMA_KEC_SUBDISTRIK':g('NAMA KEC'),'KD_DISTRIBUTOR':g('KD DISTRIBUTOR'),'DISTRIBUTOR':dist,'KD_EKSPEDITUR_ASAL':kde or None,'EKSPEDITUR_ASAL':en or None,'EXPEDITUR_KODE':ek,'EXPEDITUR_NAMA':eknm,'NOPOL':g('NOPOL'),'NO_SO':g('NO SO/NO PO'),'PRODUK':prod,'KEMASAN':kemasan_prod(prod),'KEMASAN_SUMBER':'MATERIAL','TONASE':num(g('TONASE')),'KODE_TOKO':g('KODE TOKO'),'NAMA_TOKO':g('NAMA TOKO'),'ALAMAT_TOKO':g('ALAMAT TOKO'),'NAMA_KAPAL':g('NAMA KAPAL') or None,'STATUS':g('STATUS'),'WAKTU_ANTRI':ta,'WAKTU_MASUK':tmasuk,'WAKTU_KELUAR':tk,'WAITING_TIME_JAM':wt,'DWELL_TIME_JAM':dw,'DWELL_WAITING_FLAG':fl,'DWELL_WAITING_NOTE':note,'BERAT_KOSONG':g('BERAT KOSONG'),'BERAT_ISI':g('BERAT ISI'),'BERAT_REAL':g('BERAT REAL'),'KODE_KANTONG':g('KODE KANTONG'),'ADA_DI_MASTER_LAMA':'BARU','FLAG':'; '.join(flags) or None}
    return {k:sv(v) for k,v in row.items()}
def read_rilis(path,sheets):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True);out=[];seen_sheets=[];tot=0
    for w in wb.worksheets:
        code=w.title.split(' ')[0]
        if sheets and code not in sheets: continue
        it=w.iter_rows(values_only=True);h=next(it,None)
        if not h: continue
        hm=header_map(h);need=['NO SPJ','TGL SPJ','NAMA PROPINSI','INCOTERM','TONASE']
        if 'NO SPJ' not in hm: log('  sheet %-28s dilewati (bukan sheet Rilis: kolom NO SPJ tidak ada)'%w.title);continue
        miss=[n for n in need if n not in hm]
        if miss: raise SystemExit('Sheet %s: kolom wajib tidak ditemukan: %s'%(w.title,miss))
        seen_sheets.append(code);n=0
        for r in it:
            if not r or r[hm['NO SPJ']] is None: continue
            tot+=1
            if r[hm['NAMA PROPINSI']] not in SCOPE: continue
            if r[hm['NAMA PROPINSI']] in PLANT_BATAS and code not in PLANT_BATAS[r[hm['NAMA PROPINSI']]]: continue
            rec={k:(r[i] if i<len(r) else None) for k,i in hm.items()};out.append((code,rec));n+=1
        log('  sheet %-28s baris dalam cakupan: %d'%(w.title,n))
    if not seen_sheets: raise SystemExit('Tidak ada sheet Rilis (kolom NO SPJ) di file ini. Salah slot upload?')
    return out,seen_sheets,tot
def cmd_rilis(a):
    sheets=set(a.sheets.split(',')) if a.sheets else None
    log('== RILIS',a.input,dt.datetime.now().strftime('%Y-%m-%d %H:%M'))
    recs,seen,tot=read_rilis(a.input,sheets);log('baris file (semua provinsi, sheet terbaca): %d | dalam cakupan 7 provinsi: %d'%(tot,len(recs)))
    inc_rows={};errs=[];warns=[];dup=C.Counter()
    unk=C.defaultdict(lambda:[0,0.0,set()])
    for code,rec in recs:
        row=rilis_row(rec,code);k=row['NO_SPJ'];dup[k]+=1;inc_rows[k]=row
        ton=num(row['TONASE']);d=pdate(row['TGL_SPJ'])
        if ton is None or ton<=0: errs.append('tonase tidak valid SPJ %s: %r'%(k,row['TONASE']))
        elif ton>SET['max_tonase_baris']: warns.append('tonase > %d SPJ %s: %s'%(SET['max_tonase_baris'],k,ton))
        if d is None: errs.append('tanggal tidak terbaca SPJ %s'%k)
        elif d<TMIN or d>dt.datetime.now()+dt.timedelta(days=1): errs.append('tanggal di luar rentang SPJ %s: %s'%(k,row['TGL_SPJ']))
        if code not in PLANTS: unk[code][0]+=1; unk[code][1]+=(ton or 0); unk[code][2].add(str(row['NAMA_PROPINSI']))
    for c,(n,tn,pv) in sorted(unk.items()):
        errs.append('KODE PLANT TIDAK DIKENAL: %s (%d baris, %.1f t, provinsi: %s). Tambahkan ke settings.json "plants" (kode -> nama plant), lalu upload ulang.'%(c,n,tn,', '.join(sorted(pv))))
    for k,n in dup.items():
        if n>1: errs.append('SPJ duplikat di file: %s (%dx)'%(k,n))
    store=load_store(DATA/'rilis.csv.gz',RCOLS)
    new=[];chg=[]
    for k,row in inc_rows.items():
        o=store.get(k)
        if o is None: new.append(k)
        else:
            row['ADA_DI_MASTER_LAMA']=o['ADA_DI_MASTER_LAMA']
            diff=[c for c in RCOLS if c!='ADA_DI_MASTER_LAMA' and o[c]!=row[c]]
            if diff: chg.append((k,diff,o,row))
    ds=[pdate(r['TGL_SPJ']) for r in inc_rows.values() if pdate(r['TGL_SPJ'])]
    if ds:
        lo,hi=min(ds),max(ds);cov=set(seen)
        in_range=[k for k,o in store.items() if o['KD_PLANT'] in cov and lo<=pdate(o['TGL_SPJ'])<=hi]
        missing=[k for k in in_range if k not in inc_rows]
        if in_range and len(in_range)-len(missing)<(1-SET['row_drop_pct']/100)*len(in_range) and not a.force: errs.append('guardrail: baris file < %d%% baris tersimpan pada rentang tanggal yang sama (%d dari %d)'%(100-SET['row_drop_pct'],len(in_range)-len(missing),len(in_range)))
        log('rentang tanggal file: %s s.d. %s'%(lo.strftime('%d-%m-%Y'),hi.strftime('%d-%m-%Y')))
    else: missing=[]
    # outlier harian
    byd=C.defaultdict(float)
    for r in inc_rows.values():
        d=pdate(r['TGL_SPJ'])
        if d and num(r['TONASE']): byd[(r['KD_PLANT'],d)]+=num(r['TONASE'])
    for p in {k[0] for k in byd}:
        ser=sorted((d,v) for (pp,d),v in byd.items() if pp==p)
        for i,(d,v) in enumerate(ser):
            prev=[x for _,x in ser[max(0,i-14):i]]
            if len(prev)>=5 and statistics.median(prev)>0 and v>SET['outlier_x']*statistics.median(prev) and v>1000: warns.append('lonjakan >%dx plant %s %s: %.1f t'%(SET['outlier_x'],p,d.strftime('%d-%m'),v))
    nt=C.Counter();ntt=C.defaultdict(float)
    for k in new: r=inc_rows[k];nt[(r['BULAN_NO'],r['INCOTERM'])]+=1;ntt[(r['BULAN_NO'],r['INCOTERM'])]+=num(r['TONASE']) or 0
    log('BARU: %d | BERUBAH: %d | HILANG dari file (tetap disimpan): %d | tidak berubah: %d'%(len(new),len(chg),len(missing),len(inc_rows)-len(new)-len(chg)))
    for k in sorted(nt,key=lambda x:(str(x[0]),x[1])): log('  baru bulan %s %-14s %6d baris %12.2f t'%(k[0],k[1],nt[k],ntt[k]))
    cf=C.Counter(c for _,d,_,_ in chg for c in d);log('kolom yang berubah:',dict(cf))
    for k,d,o,n in chg[:10]: log('  SPJ %s: %s'%(k,'; '.join('%s %s->%s'%(c,o[c],n[c]) for c in d[:3])))
    if missing:
        mt=sum(num(store[k]['TONASE']) or 0 for k in missing);log('  contoh hilang:',missing[:5],'total %.1f t'%mt)
    if NEW_CODES: log('KODE EKSPEDITUR BARU (perlu dicek):',NEW_CODES)
    for w in warns[:30]: log('PERINGATAN:',w)
    if len(warns)>30: log('... %d peringatan lain'%(len(warns)-30))
    for e in errs[:30]: log('ERROR:',e)
    if errs and not a.force: log('DIBATALKAN: %d error. Perbaiki sumber atau gunakan --force.'%len(errs));finish(a,'rilis');return 1
    if a.dry_run: log('DRY-RUN: tidak ada yang ditulis.');finish(a,'rilis');return 0
    for k in new: store[k]=inc_rows[k]
    for k,_,_,n in chg: store[k]=n
    save_store(DATA/'rilis.csv.gz',RCOLS,store.values())
    if NEW_CODES: json.dump(EKS,open(CFG/'ekspeditur_map.json','w'),ensure_ascii=False,indent=1)
    log('TERSIMPAN: %d baris di data/rilis.csv.gz'%len(store));finish(a,'rilis');return 0

# =============== SO ===============
SCOLS=['SRC_KEY','NO_URUT','SALES_ORDER','KODE_PRODUK','LINE','NO_OR','NO_KONTRAK','TGL_SO','TGL_KIRIM','ESTIMASI_TERIMA','BULAN_NO','PERIODE','TIPE_SO','INCOTERM','ROUTE','TOP','DISTRIBUTOR','KODE_DISTRIK','NAMA_DISTRIK','SHIP_TO','KD_PLANT','SOURCE_PLANT','PRODUK','UOM','KEMASAN','QTY_SO','QTY_RELEASE','QTY_SISA','TON_SO','TON_RELEASE','TON_SISA','NAMA_KAPAL','PROPINSI','STATUS_QTY','STATUS_HARGA','BLOCKING','REJECTION','EXPEDITUR_KODE','EXPEDITUR_NAMA','ATTRIB_METHOD','PORSI_EKSPEDITUR','ADA_DI_MASTER_LAMA','FLAG']
S_DT=['TGL_SO','TGL_KIRIM','ESTIMASI_TERIMA'];S_NUM=['BULAN_NO','QTY_SO','QTY_RELEASE','QTY_SISA','TON_SO','TON_RELEASE','TON_SISA','PORSI_EKSPEDITUR']
MON=['JANUARI','FEBRUARI','MARET','APRIL','MEI','JUNI','JULI','AGUSTUS','SEPTEMBER','OKTOBER','NOVEMBER','DESEMBER']
def target_prop():
    tp=C.defaultdict(lambda:C.defaultdict(float));p=BASE/'01_TARGET.xlsx'
    if not p.exists(): log('PERINGATAN: 01_TARGET.xlsx tidak ada; proporsi Target FRC tidak dipakai');return tp
    ws=openpyxl.load_workbook(p,read_only=True,data_only=True)['Target Ekspeditur'];it=ws.iter_rows(values_only=True);hm=header_map(next(it))
    for r in it:
        if r[hm['INCOTERM']]=='FRC' and (r[hm['TARGET_TON_BULAN']] or 0)>0: tp[(r[hm['BULAN_NO']],r[hm['KODE_DISTRIK']],str(r[hm['KODE_PLANT']]))][r[hm['EXPEDITUR_KODE']]]+=r[hm['TARGET_TON_BULAN']]
    return tp
def ton(uom,prod,q,defaults):
    if not isinstance(q,(int,float)): return 0.0
    if uom in ('TO','BAG'): return float(q)  # TO = ton; BAG = jumbo bag 1 t (diverifikasi vs Rilis: 1,000 t/BAG)
    m=re.search(r'(\d+)\s*KG',str(prod).upper())
    if not m: defaults[0]+=1;kg=50
    else: kg=int(m.group(1))
    return q*kg/1000
def cmd_so(a):
    log('== SO',a.input,dt.datetime.now().strftime('%Y-%m-%d %H:%M'))
    wb=openpyxl.load_workbook(a.input,read_only=True,data_only=True)
    if 'List SO' not in wb.sheetnames: raise SystemExit('Sheet "List SO" tidak ada di file ini (sheet: %s). Salah slot upload?'%wb.sheetnames)
    ws=wb['List SO'];it=ws.iter_rows(values_only=True);hm=header_map(next(it))
    need=['Sales Order','Product Code','Line','Tanggal Kirim','Incoterm','Qty SO','UOM','Qty Release','Qty Rest SO','Propinsi','Rejection','Blocking','Keterangan']
    miss=[n for n in need if n not in hm]
    if miss: raise SystemExit('Kolom wajib tidak ditemukan di List SO: %s'%miss)
    rows=[];total=0;dash=0
    for r in it:
        if not r or r[hm['Sales Order']] is None: continue
        total+=1
        rec={k:(r[i] if i<len(r) else None) for k,i in hm.items()}
        if str(rec.get('Keterangan') or '').strip().upper()!=SO_KET:   # SO: hanya baris Keterangan = AP1 (semua incoterm)
            dash+=1;continue
        rows.append(rec)
    log('baris file: %d | Keterangan %s: %d | bukan %s (dilewati): %d'%(total,SO_KET,len(rows),SO_KET,dash))
    full=[]
    if (DATA/'so.csv.gz').exists():
        with gzip.open(DATA/'so.csv.gz','rt',encoding='utf-8',newline='') as f: full=list(csv.DictReader(f))
    if SO_BUANG:
        n0=len(full);full=[r for r in full if r.get('PROPINSI') not in SO_BUANG]
        if len(full)!=n0: log('SO tersimpan dibuang karena bukan %s (provinsi %s): %d baris'%(SO_KET,sorted(SO_BUANG),n0-len(full)))
    for r in full:
        for c in SCOLS: r.setdefault(c,'')
        if not r['KEMASAN']: r['KEMASAN']=kemasan_uom(r.get('UOM'))
    old=C.defaultdict(list)
    for r in full: old[r['SRC_KEY']].append(r)
    frozen=json.load(open(CFG/'so_frozen_attrib_master_lama.json'));tp=target_prop()
    errs=[];warns=[];seqc=C.Counter();out={};defaults=[0];nolist=C.Counter();neg_rest=[0]
    for r in rows:
        so=str(r['Sales Order']).strip();prod=str(r['Product Code']).strip();line=r.get('Line')
        base=(so,prod,str(line));seqc[base]+=1;sk='%s|%s|%s|%d'%(so,prod,line,seqc[base])
        uom=r['UOM']
        if uom not in ('ZAK','TO','BAG'): errs.append('UOM tidak dikenal %r SO %s'%(uom,so))
        for c in ('Qty SO','Qty Release'):
            v=r[c]
            if isinstance(v,(int,float)) and v<0: errs.append('%s negatif SO %s: %s'%(c,so,v))
        if isinstance(r['Qty Rest SO'],(int,float)) and r['Qty Rest SO']<-1e-6: neg_rest[0]+=1
        tk=pdate(r['Tanggal Kirim']);bn=tk.month if tk else None;inc=r['Incoterm'];kdp=str(r.get('Plant'));kdd=r.get('District Code')
        tso,trl,tss=ton(uom,r.get('Product Name'),r['Qty SO'],defaults),ton(uom,r.get('Product Name'),r['Qty Release'],defaults),ton(uom,r.get('Product Name'),r['Qty Rest SO'],defaults)
        flags=[]
        if r['Rejection']=='YES': flags.append('REJECTION')
        if r['Blocking']=='YES': flags.append('BLOCKING')
        # atribusi
        if sk in old: spl=[(x['EXPEDITUR_KODE'],x['EXPEDITUR_NAMA'],float(x['PORSI_EKSPEDITUR'] or 0),x['ATTRIB_METHOD'],x['ADA_DI_MASTER_LAMA']) for x in old[sk]]
        elif so+'|'+prod in frozen: spl=[(e,n,f,'MASTER_LAMA:'+m,'ADA') for e,n,f,m in frozen[so+'|'+prod]]
        else:
            dist=r.get('Distributor')
            if inc!='FRC': spl=[(dist,dist,1.0,'DISTRIBUTOR (FOT/FOB/FAS)','BARU')]
            elif tk and tk<=SILOG_END: spl=[('SILOG','SEMEN INDONESIA LOGISTIK.PT',1.0,'SILOG_100_s.d._9_APR','BARU')]
            elif kdp in SET['angkut_sendiri_plants']: spl=[(dist,dist,1.0,'DISTRIBUTOR (angkut sendiri Belawan/Baturaja)','BARU')]
            else:
                z=tp.get((bn,kdd,kdp))
                if z and sum(z.values())>0:
                    T=sum(z.values());spl=[(e,e,v/T,'PROPORSI_TARGET_FRC','BARU') for e,v in z.items()]
                else: spl=[('TANPA_EKSPEDITUR','',1.0,'TANPA_EKSPEDITUR','BARU')];nolist[(bn,r.get('District Name'))]+=1
        lst=[]
        for e,n,f,m,ada in spl:
            q=lambda v:(v*f if isinstance(v,(int,float)) else None)
            row={'SRC_KEY':sk,'NO_URUT':r.get('No.'),'SALES_ORDER':so,'KODE_PRODUK':prod,'LINE':line,'NO_OR':r.get('No. OR'),'NO_KONTRAK':r.get('No. Contract'),'TGL_SO':pdate(r.get('SO Date')),'TGL_KIRIM':tk,'ESTIMASI_TERIMA':pdate(r.get('Estimasi Tgl Terima')),'BULAN_NO':bn,'PERIODE':MON[bn-1] if bn else None,'TIPE_SO':r.get('Order Type'),'INCOTERM':inc,'ROUTE':r.get('Route'),'TOP':r.get('TOP'),'DISTRIBUTOR':r.get('Distributor'),'KODE_DISTRIK':kdd,'NAMA_DISTRIK':norm_distrik(r.get('District Name')),'SHIP_TO':r.get('Ship To'),'KD_PLANT':kdp,'SOURCE_PLANT':PLANTS.get(kdp,kdp),'PRODUK':r.get('Product Name'),'KEMASAN':kemasan_uom(uom),'UOM':uom,'QTY_SO':q(r['Qty SO']),'QTY_RELEASE':q(r['Qty Release']),'QTY_SISA':q(r['Qty Rest SO']),'TON_SO':tso*f,'TON_RELEASE':trl*f,'TON_SISA':tss*f,'NAMA_KAPAL':(r.get('Ship Name') if r.get('Ship Name') not in (None,'-') else None),'PROPINSI':r['Propinsi'],'STATUS_QTY':r.get('Status Qty'),'STATUS_HARGA':r.get('Status Harga'),'BLOCKING':r['Blocking'],'REJECTION':r['Rejection'],'EXPEDITUR_KODE':e,'EXPEDITUR_NAMA':n,'ATTRIB_METHOD':m,'PORSI_EKSPEDITUR':round(f,6),'ADA_DI_MASTER_LAMA':ada,'FLAG':'; '.join(flags) or None}
            lst.append({k:sv(v) for k,v in row.items()})
        out[sk]=lst
    if neg_rest[0]: warns.append('%d baris Qty Rest SO negatif (release melebihi Qty SO; dipertahankan apa adanya)'%neg_rest[0])
    if defaults[0]: warns.append('%d konversi tonase memakai default 50 kg (nama produk tanpa "xx KG")'%defaults[0])
    newk=[k for k in out if k not in old];chg=[]
    for k,l in out.items():
        if k in old:
            a1=sum(float(x['TON_RELEASE'] or 0) for x in old[k]);b1=sum(float(x['TON_RELEASE'] or 0) for x in l)
            a2=sum(float(x['TON_SO'] or 0) for x in old[k]);b2=sum(float(x['TON_SO'] or 0) for x in l)
            f1=old[k][0]['FLAG'];f2=l[0]['FLAG'];s1=old[k][0]['STATUS_QTY'];s2=l[0]['STATUS_QTY']
            if abs(a1-b1)>0.0005 or abs(a2-b2)>0.0005 or f1!=f2 or s1!=s2: chg.append((k,a1,b1,a2,b2,f1,f2))
    missing=[k for k in old if k not in out]
    if old and len(old)-len(missing)<(1-SET['row_drop_pct']/100)*len(old) and not a.force: errs.append('guardrail: kunci file %d < 90%% kunci tersimpan %d'%(len(old)-len(missing),len(old)))
    log('SO BARU: %d | BERUBAH: %d | HILANG dari file (tetap disimpan): %d | tidak berubah: %d'%(len(newk),len(chg),len(missing),len(out)-len(newk)-len(chg)))
    nt=C.defaultdict(float);nn=C.Counter()
    for k in newk:
        for x in out[k]: nt[(x['BULAN_NO'],x['INCOTERM'])]+=float(x['TON_RELEASE'] or 0);nn[(x['BULAN_NO'],x['INCOTERM'])]+=1
    for k in sorted(nt,key=lambda x:(str(x[0]),x[1])): log('  baru bulan %s %-4s %6d baris %12.2f t release'%(k[0],k[1],nn[k],nt[k]))
    for k,a1,b1,a2,b2,f1,f2 in chg[:10]: log('  %s: release %.2f->%.2f t, SO %.2f->%.2f t, flag %s->%s'%(k,a1,b1,a2,b2,f1,f2))
    if nolist: log('TANPA_EKSPEDITUR:',sum(nolist.values()),'baris',nolist.most_common(5))
    if missing: log('  contoh SO hilang:',missing[:5])
    for w in warns[:30]: log('PERINGATAN:',w)
    for e in errs[:30]: log('ERROR:',e)
    if errs and not a.force: log('DIBATALKAN: %d error.'%len(errs));finish(a,'so');return 1
    if a.dry_run: log('DRY-RUN: tidak ada yang ditulis.');finish(a,'so');return 0
    result={}
    for k,l in old.items(): result[k]=l
    for k,l in out.items(): result[k]=l
    allrows=[x for l in result.values() for x in l]
    save_store(DATA/'so.csv.gz',SCOLS,allrows)
    tsum=sum(float(x['TON_RELEASE'] or 0) for l in out.values() for x in l);log('TERSIMPAN: %d baris (%d kunci) di data/so.csv.gz | ton release file: %.2f'%(len(allrows),len(result),tsum));finish(a,'so');return 0

def finish(a,name):
    LOGS.mkdir(exist_ok=True);p=LOGS/('update_%s_%s.txt'%(name,dt.datetime.now().strftime('%Y%m%d_%H%M%S')));p.write_text('\n'.join(LOG),encoding='utf-8');print('log:',p.name)

# =============== VIEWS ===============
def conv(r,cols,dts,nums):
    o=[]
    for c in cols:
        v=r[c]
        if v=='' or v is None: o.append(None)
        elif c in dts: o.append(pdate(v))
        elif c in nums:
            x=num(v);o.append(int(x) if x is not None and c in ('NO_SPJ','BULAN_NO') else x)
        else: o.append(v)
    return o
def cmd_rapikan(a=None):
    """Rapikan master (sumber kebenaran, berlaku juga untuk update harian): ejaan distrik, KEMASAN & KEMASAN_SUMBER di Rilis
    dicocokkan ke SO lewat NO_SO (SO unik -> kemasan SO; tanpa SO -> dari nama produk; SO campur/bertentangan -> TIDAK_JELAS)."""
    pr,ps=DATA/'rilis.csv.gz',DATA/'so.csv.gz'
    if not (pr.exists() and ps.exists()): return
    R=list(load_store(pr,RCOLS).values())
    with gzip.open(ps,'rt',encoding='utf-8',newline='') as f: S=list(csv.DictReader(f))
    so=C.defaultdict(set);ch=C.Counter()
    for r in S:
        k=r.get('KEMASAN') or kemasan_uom(r.get('UOM'));so[r['SALES_ORDER']].add(k)
        d=norm_distrik(r['NAMA_DISTRIK'])
        if d!=r['NAMA_DISTRIK']: r['NAMA_DISTRIK']=d;ch['so_distrik']+=1
    res=C.defaultdict(lambda:[0,0.0])
    for r in R:
        d=norm_distrik(r['NAMA_AREA_DISTRIK'])
        if d!=r['NAMA_AREA_DISTRIK']: r['NAMA_AREA_DISTRIK']=d;ch['rilis_distrik']+=1
        s=so.get(str(r['NO_SO']).split('.')[0]);mat=kemasan_prod(r['PRODUK'])
        if not s: src,k='MATERIAL',mat
        elif len(s)>1: src,k='TIDAK_JELAS',mat
        else:
            k=next(iter(s));src='SO'
            if k!=mat: ch['kemasan_diganti_SO']+=1
        if r['KEMASAN']!=k or r.get('KEMASAN_SUMBER')!=src: ch['kemasan_sumber']+=1
        r['KEMASAN']=k;r['KEMASAN_SUMBER']=src;res[src][0]+=1;res[src][1]+=num(r['TONASE']) or 0
    log('RAPIKAN KEMASAN Rilis:',{k:(v[0],round(v[1])) for k,v in res.items()},'| perubahan:',dict(ch))
    if ch:
        save_store(pr,RCOLS,R)
        BAK.mkdir(parents=True,exist_ok=True);shutil.copy2(ps,BAK/('so'+dt.datetime.now().strftime('_%Y%m%d_%H%M%S')+'.csv.gz'))
        with gzip.open(str(ps)+'.tmp','wt',encoding='utf-8',newline='') as f:
            w=csv.DictWriter(f,fieldnames=SCOLS);w.writeheader()
            for r in S: w.writerow({c:r.get(c,'') for c in SCOLS})
        Path(str(ps)+'.tmp').replace(ps)
def cmd_views(a):
    import xlsxwriter
    cmd_rapikan()
    for name,cols,dts,nums,fn,sheet in (('rilis',RCOLS,R_DT,R_NUM,'02_RILIS.xlsx','Rilis'),('so',SCOLS,S_DT,S_NUM,'03_SO.xlsx','SO')):
        p=DATA/(name+'.csv.gz')
        if not p.exists(): log('lewati',name,'(data belum ada)');continue
        with gzip.open(p,'rt',encoding='utf-8',newline='') as f: rows=list(csv.DictReader(f))
        wb=xlsxwriter.Workbook(str(BASE/fn),{'constant_memory':True,'default_date_format':'dd-mm-yyyy'})
        hf=wb.add_format({'bold':True,'font_color':'white','bg_color':'#3B4A63'});df=wb.add_format({'num_format':'dd-mm-yyyy'});dtf=wb.add_format({'num_format':'dd-mm-yyyy hh:mm:ss'})
        ws=wb.add_worksheet(sheet);ws.write_row(0,0,cols,hf);ws.freeze_panes(1,0)
        for i,r in enumerate(rows,1):
            for j,v in enumerate(conv(r,cols,dts,nums)):
                if v is None: continue
                if isinstance(v,dt.datetime): ws.write_datetime(i,j,v,df if cols[j] in ('TGL_SPJ','TGL_SO','TGL_KIRIM','ESTIMASI_TERIMA') else dtf)
                else: ws.write(i,j,v)
        c=wb.add_worksheet('CEK');c.write_row(0,0,['INCOTERM','BULAN_NO','BARIS','TON'],hf)
        agg=C.defaultdict(lambda:[0,0.0]);tc='TONASE' if name=='rilis' else 'TON_RELEASE'
        for r in rows: k=(r['INCOTERM'],int(r['BULAN_NO'] or 0));agg[k][0]+=1;agg[k][1]+=num(r[tc]) or 0
        for i,(k,v) in enumerate(sorted(agg.items()),1): c.write_row(i,0,[k[0],k[1],v[0],round(v[1],2)])
        if name=='rilis':
            k=wb.add_worksheet('KAMUS_EKSPEDITUR');k.write_row(0,0,['EXPEDITUR_KODE','EXPEDITUR_NAMA_LENGKAP'],hf)
            seen={}
            for r in rows: seen.setdefault(r['EXPEDITUR_KODE'],r['EXPEDITUR_NAMA'])
            for i,(c1,n1) in enumerate(sorted(seen.items()),1): k.write_row(i,0,[c1,n1])
        pf=CFG/('panduan_%s.txt'%name)
        if pf.exists():
            p2=wb.add_worksheet('PANDUAN')
            for i,l in enumerate(pf.read_text(encoding='utf-8').splitlines()): p2.write(i,0,l)
            p2.set_column(0,0,170)
        wb.close();log('view ditulis:',fn,len(rows),'baris')
    finish(a,'views')

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('cmd',choices=['rilis','so','views','rapikan']);ap.add_argument('--input');ap.add_argument('--dry-run',action='store_true');ap.add_argument('--force',action='store_true');ap.add_argument('--sheets')
    a=ap.parse_args()
    if a.cmd in ('rilis','so') and not a.input: ap.error('--input wajib')
    sys.exit({'rilis':cmd_rilis,'so':cmd_so,'views':cmd_views,'rapikan':cmd_rapikan}[a.cmd](a))
